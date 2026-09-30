"""临时冒烟：验证班级家长会 PPT 新排版（散文卡 + 编号要点卡）不溢出、不碎段。"""
import sys
from datetime import datetime

sys.path.insert(0, ".")

from app.services.pptx_builder import build_class_meeting_ppt  # noqa: E402

content = {
    "title": "Scratch 启蒙班 阶段学习汇报家长会",
    "class_summary": (
        "各位家长好！本阶段（2026年6月9日至9月9日）我们班共有2名学员在读，"
        "大家保持了非常好的学习连续性，平均出勤率达到100%，累计完成10人次·节课程。"
        "课堂上孩子们从零基础逐步掌握顺序、循环与条件三种程序结构，"
        "能够独立完成小游戏小作品的搭建，整体学习氛围积极、专注度良好。"
    ),
    "ability_comment": (
        "从班级平均来看，逻辑思维与项目实践表现突出，创意表达稳步提升；"
        "部分学员在算法抽象方面还有提升空间，下阶段将通过项目拆解练习加强。"
    ),
    "highlights": (
        "平均出勤率达到100%，学习连续性非常好；"
        "发布课后反馈8篇，记录了每位学员的课堂表现；"
        "批改作业12份，平均得分率85%；"
        "张小明、李小红等学员本阶段进步突出。"
    ),
    "to_improve": (
        "部分学员课后练习频次偏低，知识巩固不足，建议每周安排2次课后练习；"
        "课堂表达与作品讲解能力可以进一步锻炼，后续将通过课堂展示环节加强。"
    ),
    "next_plan": (
        "进入新知识模块学习，引入函数与变量概念，衔接更复杂的项目开发；"
        "每两周安排一次综合项目实践，巩固所学知识；"
        "期末组织班级作品展示会，让每位学员都有登台讲解的机会。"
    ),
    "home_suggestions": (
        "每周安排固定时间完成课后练习（建议2次，每次30分钟）；"
        "鼓励孩子把课堂作品讲给家长听，锻炼表达能力；"
        "课时不足10节时请及时与老师沟通续报，避免学习中断。"
    ),
}

out = build_class_meeting_ppt(
    class_name="Scratch 启蒙班",
    subject="Scratch 图形化编程",
    period_label="2026-06-09 ~ 2026-09-09",
    teacher_name="王老师",
    created_at=datetime(2026, 9, 10, 15, 30),
    class_stats={
        "active_count": 2,
        "student_count": 2,
        "total_lessons": 10,
        "avg_attendance_rate": 1.0,
        "total_feedbacks": 8,
        "total_homework": 12,
        "avg_homework_score_rate": 0.85,
    },
    averages=[("逻辑思维", 4.4), ("项目实践", 4.1), ("创意表达", 3.6), ("算法基础", 3.0)],
    honor_roll=[("张小明", "循环结构掌握扎实，作品完成度全班第一"), ("李小红", "课堂表达进步明显，主动帮助同学")],
    content=content,
)
print("OK ->", out)

# 校验：打开生成的 pptx，检查每页 shape 数与文本框位置不越界
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

prs = Presentation(f"uploads/{out}")
W, H = prs.slide_width, prs.slide_height
issues = []
for i, slide in enumerate(prs.slides, start=1):
    for sh in slide.shapes:
        if sh.left is None:
            continue
        if sh.left < 0 or sh.top < 0 or sh.left + sh.width > W + 9525 or sh.top + sh.height > H + 9525:
            issues.append((i, sh.shape_type, Emu(sh.left).inches, Emu(sh.top).inches))
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst), "out-of-bounds:", issues)
